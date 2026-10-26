from pathlib import Path
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter
from analyze_evening_pickup import load_matched_orders

ROOT = Path(os.environ.get("TMALL_LOGISTICS_ROOT", Path(__file__).resolve().parent))
BOOK = Path(os.environ["TMALL_LOGISTICS_OUTPUT"])
N = 0
RED = "C72C41"
DARK = "252525"
GREY = "F3F4F6"
LIGHT = "FCEEEF"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D9D9D9")


def parse_times(value):
    return [tuple(map(int, part.split(":"))) for part in value.split(",")]


CURRENT_TIMES = parse_times(os.environ.get("TMALL_LOGISTICS_PICKUPS", "10:00,17:00"))
TIMES = parse_times(os.environ.get("TMALL_LOGISTICS_EXTRA_PICKUPS", "17:30,18:00,18:30,19:00,19:30,20:00,20:30,21:00,21:30,22:00,22:30,23:00"))


def time_text(hour, minute):
    return f"{hour:02d}:{minute:02d}"


def pickup_expression(row, times):
    ordered = sorted(times)
    first_hour, first_minute = ordered[0]
    expression = f"INT(H{row})+1+TIME({first_hour},{first_minute},0)"
    for hour, minute in reversed(ordered):
        expression = f"IF(H{row}<=INT(H{row})+TIME({hour},{minute},0),INT(H{row})+TIME({hour},{minute},0),{expression})"
    return expression


def setup_sheet(wb, name):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    return ws


def title(ws, text, end_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, text)
    cell.font = Font(name="PingFang SC", size=18, bold=True, color=DARK)
    cell.fill = PatternFill("solid", fgColor=LIGHT)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38


def header(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(name="PingFang SC", color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[row].height = 30


def score_formula(cell):
    return f'=IF({cell}<=43.99,"5.0档",IF({cell}<=47.36,"4.8档",IF({cell}<=49.91,"4.6档",IF({cell}<=56,"4.2档",IF({cell}<=58.66,"4.0档","未达4.0档")))))'


def add_detail_sheet(wb, matched):
    ws = setup_sheet(wb, "晚间揽收订单明细")
    ws.freeze_panes = "A2"
    scenario_headers = [f"{time_text(h, m)}节省/小时" for h, m in TIMES]
    columns = [
        "到货明细订单编号", "匹配子订单编号", "主订单编号", "商品SKU", "物流线路", "快递公司",
        "订单付款时间", "发货时间", "原到货时长/小时", "匹配方式", "适用本店班次",
        "当前模拟揽收", "当前等待/小时", "发货小时",
    ] + scenario_headers
    ws.append(columns)
    for i, row in matched.iterrows():
        excel_row = i + 2
        ws.append([
            str(row["订单编号"]), str(row["子订单编号"]), str(row["主订单编号"]), row["商品sku"],
            row["物流线路"], row["快递公司"], row["订单付款时间"].to_pydatetime(),
            row["发货时间"].to_pydatetime(), float(row["物流到货时长/小时"]), row["匹配方式"],
        ])
        origin = os.environ.get("TMALL_LOGISTICS_ORIGIN", "深圳市")
        courier = os.environ.get("TMALL_LOGISTICS_COURIER", "韵达快递")
        ws.cell(excel_row, 11, f'=IF(AND(LEFT(E{excel_row},3)="{origin}",F{excel_row}="{courier}",H{excel_row}<>""),1,0)')
        ws.cell(excel_row, 12, f'={pickup_expression(excel_row, CURRENT_TIMES)}')
        ws.cell(excel_row, 13, f'=(L{excel_row}-H{excel_row})*24')
        ws.cell(excel_row, 14, f'=HOUR(H{excel_row})')
        for scenario_index, extra_time in enumerate(TIMES, start=15):
            new_pickup = pickup_expression(excel_row, CURRENT_TIMES + [extra_time])
            ws.cell(excel_row, scenario_index, f'=IF($K{excel_row}=1,MAX(0,($L{excel_row}-({new_pickup}))*24),0)')
        for col in (1, 2, 3):
            ws.cell(excel_row, col).number_format = "@"
        for col in (7, 8, 12):
            ws.cell(excel_row, col).number_format = "yyyy-mm-dd hh:mm:ss"
        for col in [9, 13] + list(range(15, 15 + len(TIMES))):
            ws.cell(excel_row, col).number_format = "0.00"
    header(ws, 1, len(columns))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{N + 1}"
    widths = {1: 22, 2: 22, 3: 22, 4: 44, 5: 24, 6: 14, 7: 21, 8: 21, 9: 18, 10: 16, 11: 16, 12: 21, 13: 16, 14: 12}
    for col in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 17)
    return ws


def add_scenario_sheet(wb):
    ws = setup_sheet(wb, "晚间揽收分析")
    ws.freeze_panes = "A7"
    title(ws, "物流到货时长考核订单：新增晚间揽收时点测算", 10)
    notes = [
        "口径：以平台导出的考核明细为准；仅当前发货城市＋快递组合匹配的订单适用本次揽收班次调整。",
        "假设：后台‘发货时间’视为包裹可交接时间；新增揽收后，后续干线与签收按相同小时数提前。实际揽收扫描未导出。",
        "因此下表是同步提前情景，不是效果保证；若只有揽收扫描提前、干线仍次日发运，实际收益可能为零或低于模拟。",
    ]
    for row, text in enumerate(notes, 2):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        ws.cell(row, 1, text)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 25
    best_row = 7 + [time_text(h, m) for h, m in TIMES].index("18:00")
    columns = ["新增揽收时间", "受益订单数", "受益订单占比", "受益单平均节省/小时", "总节省/小时", "全店平均缩短/小时", "模拟平均到货/小时", "达到档位", "较18点少节省/小时", "执行观察"]
    for col, value in enumerate(columns, 1):
        ws.cell(6, col, value)
    header(ws, 6, len(columns))
    for index, (hour, minute) in enumerate(TIMES):
        row = 7 + index
        detail_col = get_column_letter(15 + index)
        source = f"'晚间揽收订单明细'!${detail_col}$2:${detail_col}${N + 1}"
        ws.cell(row, 1, time_text(hour, minute))
        ws.cell(row, 2, f'=COUNTIF({source},">0")')
        ws.cell(row, 3, f'=B{row}/{N}')
        ws.cell(row, 4, f'=AVERAGEIF({source},">0")')
        ws.cell(row, 5, f'=SUM({source})')
        ws.cell(row, 6, f'=E{row}/{N}')
        ws.cell(row, 7, f'=AVERAGE(订单明细!$F$2:$F${N + 1})-F{row}')
        ws.cell(row, 8, score_formula(f"G{row}"))
        ws.cell(row, 9, f'=$E${best_row}-E{row}')
        ws.cell(row, 10, f'=IF(A{row}="18:00","推荐",IF(B{row}<$B${best_row},"覆盖不足","覆盖相同，节省递减"))')
        ws.cell(row, 3).number_format = "0.00%"
        for col in (4, 5, 6, 7, 9):
            ws.cell(row, col).number_format = "0.000"
        if row == best_row:
            for col in range(1, 11):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=LIGHT)
                ws.cell(row, col).font = Font(name="PingFang SC", bold=True, color=RED)
    ws.auto_filter.ref = f"A6:J{6 + len(TIMES)}"
    ws.conditional_formatting.add(f"E7:E{6 + len(TIMES)}", DataBarRule(start_type="num", start_value=0, end_type="max", color=RED))

    sensitivity_title_row = 21
    ws.cell(sensitivity_title_row, 1, "18:00 情景：实际兑现比例敏感性")
    ws.cell(sensitivity_title_row, 1).font = Font(name="PingFang SC", size=14, bold=True, color=RED)
    sensitivity_headers = ["实际兑现比例", "全店平均缩短/小时", "模拟平均到货/小时", "达到档位", "解释"]
    for col, value in enumerate(sensitivity_headers, 1):
        ws.cell(22, col, value)
    header(ws, 22, len(sensitivity_headers))
    for row, ratio in enumerate([0, 0.25, 0.5, 0.75, 1], 23):
        ws.cell(row, 1, ratio)
        ws.cell(row, 2, f'=$F${best_row}*A{row}')
        ws.cell(row, 3, f'=AVERAGE(订单明细!$F$2:$F${N + 1})-B{row}')
        ws.cell(row, 4, score_formula(f"C{row}"))
        ws.cell(row, 5, "仅用于检验实际效果低于同步提前情景时的档位变化")
        ws.cell(row, 1).number_format = "0%"
        ws.cell(row, 2).number_format = "0.000"
        ws.cell(row, 3).number_format = "0.000"
    for col, width in enumerate([18, 16, 18, 23, 18, 22, 22, 18, 22, 28], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    return ws


def update_summary(wb, matched, scenarios):
    from analyze_evening_pickup import run_scenarios
    _, scenario_results = run_scenarios(matched)
    best_row = 7 + [time_text(h, m) for h, m in TIMES].index("18:00")
    best = scenario_results.loc[scenario_results["新增揽收时间"] == "18:00"].iloc[0]
    direct_count = int((matched["匹配方式"] == "子订单号").sum())
    fallback_count = len(matched) - direct_count
    applicable = int(matched["适用本店班次"].sum())
    ws = wb["分析摘要"]
    ws.cell(32, 1, "晚间揽收增补结论")
    ws.cell(32, 1).font = Font(name="PingFang SC", size=14, bold=True, color=RED)
    metrics = [
        ("指标", "结果", "说明"),
        ("完整考核记录", f"=COUNTA('晚间揽收订单明细'!A2:A{N + 1})", f"{direct_count}笔子订单号直配；{fallback_count}笔主订单号＋SKU匹配"),
        ("适用本店班次订单", f"=SUM('晚间揽收订单明细'!K2:K{ N + 1})", f"{applicable}笔适用当前班次调整"),
        ("18:00受益订单数", f"='晚间揽收分析'!B{best_row}", "以发货时间作为可交接时间"),
        ("18:00全店平均缩短/小时", f"='晚间揽收分析'!F{best_row}", "同步提前情景"),
        ("18:00模拟平均到货/小时", f"='晚间揽收分析'!G{best_row}", f"当前为{matched['物流到货时长/小时'].mean():.3f}小时"),
        ("18:00模拟达到档位", f"='晚间揽收分析'!H{best_row}", "按当前页面档位线"),
        ("距4.2档仍差/小时", f"=MAX('晚间揽收分析'!G{best_row}-56,0)", "需叠加异常件治理或线路提速"),
    ]
    for r, values in enumerate(metrics, 33):
        for c, value in enumerate(values, 1):
            ws.cell(r, c, value)
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="center")
    header(ws, 33, 3)
    for r in range(34, 41):
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=GREY)
        ws.cell(r, 1).font = Font(name="PingFang SC", bold=True)
        ws.row_dimensions[r].height = 24
    for r in (37, 38, 40):
        ws.cell(r, 2).number_format = "0.000"
    ws.merge_cells("A42:F42")
    ws["A42"] = f"18:00情景覆盖{int(best['受益订单数'])}笔，模拟全店缩短{best['全店平均缩短小时']:.3f}小时。先试行7天并记录真实揽收、干线发运和签收时间。"
    ws["A42"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[42].height = 34
    ws.merge_cells("A43:F43")
    ws["A43"] = "注意：以上是基于同步提前假设的情景测算。若晚间仅提前揽收扫描、干线仍次日发运，则不能据此承诺提分。"
    ws["A43"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[43].height = 30


def update_sources(wb, matched):
    ws = wb["计分规则与来源"]
    ws.cell(22, 1, "晚间揽收增补：数据来源与假设")
    ws.cell(22, 1).font = Font(name="PingFang SC", size=14, bold=True, color=RED)
    direct_count = int((matched["匹配方式"] == "子订单号").sum())
    fallback_count = len(matched) - direct_count
    mean_hours = matched["物流到货时长/小时"].mean()
    normal_count = int((matched["物流到货时长/小时"] <= matched["线路标准时长/小时"]).sum())
    rows = [
        ("项目", "规则/数值", "来源", "备注"),
        ("当前揽收班次", os.environ.get("TMALL_LOGISTICS_PICKUPS", "10:00,17:00"), "运行参数", "用于计算当前下一班次"),
        ("子订单时间数据", Path(os.environ["TMALL_LOGISTICS_ITEMS"]).name, "千牛宝贝销售明细报表", f"按{N}笔考核记录匹配"),
        ("匹配完整性", f"{N}/{N}笔；发货时间{N}/{N}非空", "本次数据核对", f"{direct_count}笔子订单号直配；{fallback_count}笔主订单号＋SKU匹配"),
        ("考核范围核验", f"{N}笔均值{mean_hours:.3f}小时；需与后台展示值核对", "体验分明细与导出数据", f"{normal_count}笔未超线路标准；不是低分异常订单精选"),
        ("可交接时间代理", "商家后台发货时间", "分析假设", "不是实际快递揽收扫描时间"),
        ("同步提前情景", "新到货时长＝原到货时长－减少的交接等待", "分析假设", "后续运输与签收等幅提前；实际效果需试跑验证"),
        ("未兑现风险", "仅扫描提前、干线不发运时可能无实际收益", "分析边界", "不可把模拟结果视为平台提分保证"),
    ]
    for r, values in enumerate(rows, 23):
        for c, value in enumerate(values, 1):
            ws.cell(r, c, value)
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    header(ws, 23, 4)


def main():
    global N
    matched = load_matched_orders()
    N = len(matched)
    if "18:00" not in [time_text(h, m) for h, m in TIMES]:
        raise ValueError("晚间候选时间必须包含18:00，便于汇总与敏感性分析")
    wb = load_workbook(BOOK)
    add_detail_sheet(wb, matched)
    add_scenario_sheet(wb)
    wb["晚间揽收分析"]["A2"] = f"口径：{N}笔考核记录全部匹配；其中{int(matched['适用本店班次'].sum())}笔适用本店班次调整。"
    update_summary(wb, matched, None)
    update_sources(wb, matched)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(BOOK)
    print(BOOK)


if __name__ == "__main__":
    main()
