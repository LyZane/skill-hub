import os
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import DataBarRule, ColorScaleRule
from openpyxl.utils import get_column_letter

root=Path(os.environ.get('TMALL_LOGISTICS_ROOT', Path(__file__).resolve().parent))
src=Path(os.environ['TMALL_LOGISTICS_ARRIVALS'])
out=Path(os.environ['TMALL_LOGISTICS_OUTPUT'])
df=pd.read_excel(src,dtype={'订单编号':str})
df['物流线路']=df['物流线路'].astype(str).str.replace(r'\s*-\s*',' - ',regex=True)
n=len(df)
wb=Workbook()
wb.remove(wb.active)
red='C72C41'; dark='252525'; grey='F3F4F6'; light='FCEEEF'; white='FFFFFF'
thin=Side(style='thin',color='D9D9D9')

def base_sheet(name):
    ws=wb.create_sheet(name)
    ws.sheet_view.showGridLines=False
    ws.freeze_panes='A2'
    return ws

def header(ws,row,cols):
    for c in range(1,cols+1):
        cell=ws.cell(row,c)
        cell.fill=PatternFill('solid',fgColor=dark)
        cell.font=Font(color=white,bold=True)
        cell.alignment=Alignment(horizontal='center',vertical='center')
        cell.border=Border(bottom=thin)
    ws.row_dimensions[row].height=24

def title(ws,text,end_col):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=end_col)
    c=ws.cell(1,1,text)
    c.font=Font(size=18,bold=True,color=dark)
    c.fill=PatternFill('solid',fgColor=light)
    c.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=38

# 订单明细
ws=base_sheet('订单明细')
cols=['订单编号','商品编码/名称','商品sku','物流线路','快递公司','物流到货时长/小时','线路标准时长/小时','发货城市','收货城市','超标准时长/小时','超标准正数/小时','压至线路标准模拟','深圳韵达48小时模拟','深圳韵达36小时模拟','深圳韵达24小时模拟','四城48小时模拟','组合方案模拟']
ws.append(cols)
for i,r in df.iterrows():
    row=i+2
    vals=[str(r['订单编号']),r['商品编码/名称'],r['商品sku'],r['物流线路'],r['快递公司'],float(r['物流到货时长/小时']),float(r['线路标准时长/小时'])]
    ws.append(vals)
    ws.cell(row,8,f'=TRIM(LEFT(D{row},FIND("-",D{row})-1))')
    ws.cell(row,9,f'=TRIM(MID(D{row},FIND("-",D{row})+1,50))')
    ws.cell(row,10,f'=F{row}-G{row}')
    ws.cell(row,11,f'=MAX(J{row},0)')
    ws.cell(row,12,f'=MIN(F{row},G{row})')
    local=f'AND($D{row}="深圳市 - 深圳市",$E{row}="韵达快递")'
    ws.cell(row,13,f'=IF({local},MIN(F{row},48),F{row})')
    ws.cell(row,14,f'=IF({local},MIN(F{row},36),F{row})')
    ws.cell(row,15,f'=IF({local},MIN(F{row},24),F{row})')
    ws.cell(row,16,f'=IF(OR($I{row}="青岛市",$I{row}="荆门市",$I{row}="德州市",$I{row}="邢台市"),MIN(F{row},48),F{row})')
    ws.cell(row,17,f'=IF({local},MIN(F{row},36),MIN(F{row},G{row}))')
    ws.cell(row,1).number_format='@'
header(ws,1,len(cols)); ws.auto_filter.ref=f'A1:Q{n+1}'
widths=[22,45,36,24,14,20,22,14,18,20,20,20,22,22,22,20,20]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
for row in ws.iter_rows(min_row=2,max_row=n+1,min_col=6,max_col=17):
    for c in row: c.number_format='0.00'
ws.conditional_formatting.add(f'J2:J{n+1}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFFFFF',end_type='max',end_color='F8696B'))

# 城市分析
cities=(df.assign(收货城市=df['物流线路'].str.split(' - ').str[1],超=df['物流到货时长/小时']-df['线路标准时长/小时']).groupby('收货城市').agg(订单数=('订单编号','size'),超总=('超',lambda x:x.clip(lower=0).sum())).reset_index().sort_values(['超总','订单数'],ascending=False))
ws=base_sheet('城市分析')
cols=['收货城市','订单数','订单占比','平均到货时长/小时','平均标准时长/小时','超标准单数','超标准总小时','压至标准全店缩短/小时','压至标准后全店均值/小时','每单提速24小时的全店缩短/小时','快速专线目标/小时','快速专线全店缩短/小时','快速专线后全店均值/小时']
ws.append(cols)
for i,city in enumerate(cities['收货城市'],2):
    ws.cell(i,1,city)
    ws.cell(i,2,f'=COUNTIF(订单明细!$I$2:$I${n+1},A{i})')
    ws.cell(i,3,f'=B{i}/{n}')
    ws.cell(i,4,f'=AVERAGEIF(订单明细!$I$2:$I${n+1},A{i},订单明细!$F$2:$F${n+1})')
    ws.cell(i,5,f'=AVERAGEIF(订单明细!$I$2:$I${n+1},A{i},订单明细!$G$2:$G${n+1})')
    ws.cell(i,6,f'=COUNTIFS(订单明细!$I$2:$I${n+1},A{i},订单明细!$J$2:$J${n+1},">0")')
    ws.cell(i,7,f'=SUMIF(订单明细!$I$2:$I${n+1},A{i},订单明细!$K$2:$K${n+1})')
    ws.cell(i,8,f'=G{i}/{n}')
    ws.cell(i,9,f'=AVERAGE(订单明细!$F$2:$F${n+1})-H{i}')
    ws.cell(i,10,f'=B{i}*24/{n}')
    ws.cell(i,11,f'=IF(A{i}="深圳市",24,48)')
    ws.cell(i,12,f'=SUMPRODUCT((订单明细!$I$2:$I${n+1}=A{i})*(订单明细!$F$2:$F${n+1}>K{i})*(订单明细!$F$2:$F${n+1}-K{i}))/{n}')
    ws.cell(i,13,f'=AVERAGE(订单明细!$F$2:$F${n+1})-L{i}')
header(ws,1,len(cols)); ws.auto_filter.ref=f'A1:M{len(cities)+1}'
for i,w in enumerate([20,12,14,22,22,16,18,26,28,32,22,26,28],1): ws.column_dimensions[get_column_letter(i)].width=w
for row in ws.iter_rows(min_row=2,max_row=len(cities)+1,min_col=3,max_col=13):
    for c in row: c.number_format='0.00%' if c.column==3 else '0.000'
ws.conditional_formatting.add(f'G2:G{len(cities)+1}',DataBarRule(start_type='num',start_value=0,end_type='max',color=red))

# 线路分析
routes=(df.assign(超=df['物流到货时长/小时']-df['线路标准时长/小时']).groupby(['物流线路','快递公司']).agg(订单数=('订单编号','size'),超总=('超',lambda x:x.clip(lower=0).sum())).reset_index().sort_values(['超总','订单数'],ascending=False))
ws=base_sheet('线路分析')
cols=['物流线路','快递公司','订单数','订单占比','平均到货时长/小时','平均标准时长/小时','超标准单数','超标准总小时','压至标准全店缩短/小时','压至标准后全店均值/小时']
ws.append(cols)
for i,r in routes.iterrows():
    row=i+2
    ws.cell(row,1,r['物流线路']); ws.cell(row,2,r['快递公司'])
    ws.cell(row,3,f'=COUNTIFS(订单明细!$D$2:$D${n+1},A{row},订单明细!$E$2:$E${n+1},B{row})')
    ws.cell(row,4,f'=C{row}/{n}')
    ws.cell(row,5,f'=AVERAGEIFS(订单明细!$F$2:$F${n+1},订单明细!$D$2:$D${n+1},A{row},订单明细!$E$2:$E${n+1},B{row})')
    ws.cell(row,6,f'=AVERAGEIFS(订单明细!$G$2:$G${n+1},订单明细!$D$2:$D${n+1},A{row},订单明细!$E$2:$E${n+1},B{row})')
    ws.cell(row,7,f'=COUNTIFS(订单明细!$D$2:$D${n+1},A{row},订单明细!$E$2:$E${n+1},B{row},订单明细!$J$2:$J${n+1},">0")')
    ws.cell(row,8,f'=SUMIFS(订单明细!$K$2:$K${n+1},订单明细!$D$2:$D${n+1},A{row},订单明细!$E$2:$E${n+1},B{row})')
    ws.cell(row,9,f'=H{row}/{n}')
    ws.cell(row,10,f'=AVERAGE(订单明细!$F$2:$F${n+1})-I{row}')
header(ws,1,len(cols)); ws.auto_filter.ref=f'A1:J{len(routes)+1}'
for i,w in enumerate([26,16,12,14,22,22,16,18,26,28],1): ws.column_dimensions[get_column_letter(i)].width=w
for row in ws.iter_rows(min_row=2,max_row=len(routes)+1,min_col=4,max_col=10):
    for c in row: c.number_format='0.00%' if c.column==4 else '0.000'
ws.conditional_formatting.add(f'H2:H{len(routes)+1}',DataBarRule(start_type='num',start_value=0,end_type='max',color=red))

# 情景测算
ws=base_sheet('情景测算')
ws.append(['方案','模拟后全店平均时长/小时','较当前缩短/小时','达到到货时长档位','方案解释'])
over_count=int((df['物流到货时长/小时']>df['线路标准时长/小时']).sum())
local_mask=df['物流线路'].eq('深圳市 - 深圳市')&df['快递公司'].eq('韵达快递')
local_count=int(local_mask.sum())
scenarios=[
('当前基准','F',f'原始{n}笔考核订单'),
('全部超标准单压至各自线路标准','L',f'只治理{over_count}笔超标准订单，其余不变'),
('深圳市内韵达全部压至48小时','M',f'{local_count}笔深圳市内韵达订单设置48小时上限'),
('深圳市内韵达全部压至36小时','N',f'{local_count}笔深圳市内韵达订单设置36小时上限'),
('深圳市内韵达全部压至24小时','O',f'{local_count}笔深圳市内韵达订单设置24小时上限'),
('青岛/荆门/德州/邢台全部压至48小时','P','四城订单设置48小时上限，其余不变'),
('深圳韵达36小时+其余超标准单压至标准','Q','组合治理方案')]
for i,(name,col,note) in enumerate(scenarios,2):
    ws.cell(i,1,name)
    ws.cell(i,2,f'=AVERAGE(订单明细!${col}$2:${col}${n+1})')
    ws.cell(i,3,f'=$B$2-B{i}')
    ws.cell(i,4,f'=IF(B{i}<=43.99,"5.0档",IF(B{i}<=47.36,"4.8档",IF(B{i}<=49.91,"4.6档",IF(B{i}<=56,"4.2档",IF(B{i}<=58.66,"4.0档","未达4.0档")))))')
    ws.cell(i,5,note)
header(ws,1,5)
for i,w in enumerate([48,30,24,22,56],1): ws.column_dimensions[get_column_letter(i)].width=w
for row in ws.iter_rows(min_row=2,max_row=8,min_col=2,max_col=3):
    for c in row: c.number_format='0.000'
for r in range(2,9): ws.cell(r,5).alignment=Alignment(wrap_text=True,vertical='top')

# 摘要
ws=base_sheet('分析摘要'); ws.freeze_panes='A4'; title(ws,'物流到货时长：城市与专线收益分析',6)
period=os.environ.get('TMALL_LOGISTICS_PERIOD','以平台导出明细为准；不是最近30天全量销售')
summary=[('数据周期',period),('考核订单数',f'=COUNTA(订单明细!A2:A{n+1})'),('当前平均到货时长/小时',f'=AVERAGE(订单明细!F2:F{n+1})'),('超标准订单数',f'=COUNTIF(订单明细!J2:J{n+1},">0")'),('超标准订单占比','=B6/B4'),('全部超标准正数小时',f'=SUM(订单明细!K2:K{n+1})'),('当前到货时长得分',float(os.environ.get('TMALL_LOGISTICS_ARRIVAL_SCORE','3.9'))),('48小时揽收及时率得分',float(os.environ.get('TMALL_LOGISTICS_PICKUP_SCORE','4.9'))),('当前物流速度分',float(os.environ.get('TMALL_LOGISTICS_SPEED_SCORE','4.5'))),('当前真实体验分',float(os.environ.get('TMALL_LOGISTICS_OVERALL_SCORE','4.76')))]
for i,(k,v) in enumerate(summary,3):
    ws.cell(i,1,k); ws.cell(i,2,v)
    ws.cell(i,1).font=Font(bold=True,color=dark)
    ws.cell(i,1).fill=PatternFill('solid',fgColor=grey)
for r in [5,8,9,10,11,12]: ws.cell(r,2).number_format='0.000'
ws.cell(7,2).number_format='0.00%'
ws.cell(15,1,'关键结论'); ws.cell(15,1).font=Font(size=14,bold=True,color=red)
actual=df['物流到货时长/小时']; standard=df['线路标准时长/小时']; excess=(actual-standard).clip(lower=0)
over_count=int((actual>standard).sum()); all_standard_mean=(actual-excess).mean()
local_mask=df['物流线路'].eq('深圳市 - 深圳市')&df['快递公司'].eq('韵达快递')
local_count=int(local_mask.sum()); local_over=int((actual[local_mask]>standard[local_mask]).sum()); local_excess=excess[local_mask].sum()
local_share=local_excess/excess.sum() if excess.sum() else 0
conclusions=[
f'1. {n}笔订单中{over_count}笔超过线路标准；异常订单与线路的贡献需结合单量判断。',
f'2. 深圳市内韵达{local_count}笔，其中{local_over}笔超标准，贡献{local_excess:.2f}小时超额，占全部超额{local_share:.2%}。',
'3. 城市优先级应同时看超标准总小时、订单量和专线可兑现目标，不宜只看城市平均值。',
'4. 单量小且仅个别异常的城市，先处理异常件，不建议仅凭单笔异常开专线。',
f'5. 所有超标准单压至线路标准后，全店均值为{all_standard_mean:.3f}小时；再对照当前档位线判断后续空间。']
for i,t in enumerate(conclusions,16):
    ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=6)
    ws.cell(i,1,t); ws.cell(i,1).alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[i].height=30
ws.cell(23,1,'重点方案测算'); ws.cell(23,1).font=Font(size=14,bold=True,color=red)
ws.append([])
for j,v in enumerate(['方案','模拟后平均时长/小时','全店缩短/小时','达到档位'],1): ws.cell(24,j,v)
for idx,srcrow in enumerate([3,4,5,6,7,8],25):
    ws.cell(idx,1,f'=情景测算!A{srcrow}'); ws.cell(idx,2,f'=情景测算!B{srcrow}'); ws.cell(idx,3,f'=情景测算!C{srcrow}'); ws.cell(idx,4,f'=情景测算!D{srcrow}')
    ws.cell(idx,2).number_format='0.000'; ws.cell(idx,3).number_format='0.000'
header(ws,24,4)
for i,w in enumerate([52,26,24,22,16,16],1): ws.column_dimensions[get_column_letter(i)].width=w

# 计分规则与来源
ws=base_sheet('计分规则与来源'); title(ws,'计分规则、统计口径与数据来源',4)
rows=[
('项目','规则/数值','来源','备注'),
('物流到货时长','支付至签收总时长 ÷ 支付且已签收子订单数','千牛《店铺真实体验分运营手册》','考核周期T-40至T-11；不是最近30天全量销售'),
('考核订单范围',period,f'商家后台{n}笔订单明细',f'{n}笔均值{df["物流到货时长/小时"].mean():.3f}小时；运行时必须与后台展示值核对'),
('物流速度权重','真实体验分的20%','商家后台当前页面','运行前确认类目和权重是否变化'),
('48小时揽收及时率权重','物流速度分的60%','商家后台当前页面','当前脚本默认指标得分4.9，可通过参数调整'),
('物流到货时长权重','物流速度分的40%','商家后台当前页面','当前脚本默认指标得分3.9，可通过参数调整'),
('真实体验分公式','宝贝质量×35%+物流速度×20%+服务保障×45%','商家后台当前权重','临时情景计算时保持其他指标不变'),
('档位计分方式','按所处类目的指标档位线计算，具体以产品页面为准','千牛官方运营手册','本报告以是否达到档位为主，不虚构中间分公式'),
('官方规则URL','https://qn.taobao.com/headline/news/10721647','千牛工作台','店铺真实体验分运营手册'),
('原始数据',src.name,'商家后台导出',f'{n}笔物流到货时长订单明细')]
for r in rows: ws.append(r)
header(ws,2,4)
for i,w in enumerate([28,68,36,54],1): ws.column_dimensions[get_column_letter(i)].width=w
for row in ws.iter_rows(min_row=3,max_row=11):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
ws.append([]); ws.append(['到货时长得分档位','时长上限/小时','',''])
header(ws,13,2)
for score,hours in [(3.0,72.00),(4.0,58.66),(4.2,56.00),(4.6,49.91),(4.8,47.36),(5.0,43.99)]: ws.append([score,hours])

# 全局格式
for ws in wb.worksheets:
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1
    ws.sheet_view.zoomScale=90
    for row in ws.iter_rows():
        for c in row:
            c.font=Font(name='PingFang SC',size=c.font.sz or 11,bold=c.font.bold,color=c.font.color)
            if c.row>1 and c.alignment==Alignment(): c.alignment=Alignment(vertical='center')
wb.calculation.fullCalcOnLoad=True
wb.calculation.forceFullCalc=True
wb.calculation.calcMode='auto'
wb.save(out)
print(out)
