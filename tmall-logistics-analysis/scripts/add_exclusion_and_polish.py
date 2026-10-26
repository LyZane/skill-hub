from pathlib import Path
import math
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(os.environ.get("TMALL_LOGISTICS_ROOT", Path(__file__).resolve().parent))
BOOK = Path(os.environ["TMALL_LOGISTICS_OUTPUT"])
ARRIVALS = Path(os.environ["TMALL_LOGISTICS_ARRIVALS"])
DARK = "252525"
RED = "C72C41"
LIGHT = "FCEEEF"
GREY = "F5F6F7"
MID_GREY = "E7E9EB"
WHITE = "FFFFFF"
BLUE = "3D7FFF"
THIN = Side(style="thin", color=MID_GREY)


def font(bold=False, color=DARK, size=10):
    return Font(name="PingFang SC", size=size, bold=bold, color=color)


def add_exclusion_sheet(wb):
    df = pd.read_excel(ARRIVALS, dtype={"订单编号": str})
    df["商品ID"] = df["商品编码/名称"].str.extract(r"/(\d+)\s*$")
    df["收货城市"] = df["物流线路"].str.split("-").str[1].str.strip()
    product_id = os.environ["TMALL_LOGISTICS_EXCLUDE_PRODUCT_ID"]
    city = os.environ["TMALL_LOGISTICS_EXCLUDE_CITY"]
    mask = df["商品ID"].eq(product_id) & df["收货城市"].eq(city)
    excluded = df.loc[mask]
    n = len(df)
    if excluded.empty:
        raise ValueError("临时剔除条件没有匹配订单，未生成该情景")

    orders = wb["订单明细"]
    orders.cell(1, 18, "商品ID")
    for row in range(2, n + 2):
        orders.cell(row, 18, f'=TRIM(RIGHT(SUBSTITUTE(B{row},"/",REPT(" ",30)),30))')
        orders.cell(row, 18).number_format = "@"
    orders.auto_filter.ref = f"A1:R{n + 1}"
    orders.column_dimensions["R"].width = 18

    name = "临时剔除测算"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.merge_cells("A1:F1")
    ws["A1"] = f"临时情景：剔除{city}收货 × 商品 {product_id}"
    ws["A1"].font = font(True, RED, 18)
    ws["A1"].fill = PatternFill("solid", fgColor=LIGHT)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A2:F2")
    ws["A2"] = "仅用于假设核算，不计入标准改善方案；不代表平台允许剔除这些订单。"
    ws["A2"].font = font(True, RED, 10)
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 26

    rows = [
        ("指标", "结果", "口径/说明"),
        ("剔除商品ID", product_id, "用户指定临时条件"),
        ("剔除收货城市", city, "物流线路的收货城市，不是发货城市"),
        ("原考核订单数", f"=COUNTA(订单明细!A2:A{n + 1})", "平台物流到货时长考核明细"),
        ("原平均到货/小时", f"=AVERAGE(订单明细!F2:F{n + 1})", "剔除前"),
        ("剔除订单数", f'=COUNTIFS(订单明细!$I$2:$I${n + 1},$B$5,订单明细!$R$2:$R${n + 1},$B$4)', "同时满足城市与商品ID"),
        ("剔除订单总时长/小时", f'=SUMIFS(订单明细!$F$2:$F${n + 1},订单明细!$I$2:$I${n + 1},$B$5,订单明细!$R$2:$R${n + 1},$B$4)', "被剔除订单的到货时长合计"),
        ("剔除订单平均/小时", "=B9/B8", "用于判断剔除样本偏快还是偏慢"),
        ("剩余订单数", "=B6-B8", "原考核订单数－剔除订单数"),
        ("剔除后平均到货/小时", f"=(SUM(订单明细!$F$2:$F${n + 1})-B9)/B11", "剩余订单重新计算平均值"),
        ("均值变化/小时", "=B12-B7", "正数表示变差，负数表示改善"),
        ("到货时长档位", '=IF(B12<=43.99,"5.0档",IF(B12<=47.36,"4.8档",IF(B12<=49.91,"4.6档",IF(B12<=56,"4.2档",IF(B12<=58.66,"4.0档","未达4.0档")))))', "按当前页面档位线"),
        ("模拟物流速度分", '=4.9*60%+IF(B12<=43.99,5,IF(B12<=47.36,4.8,IF(B12<=49.91,4.6,IF(B12<=56,4.2,IF(B12<=58.66,4,3.9)))))*40%', "其他指标不变；未达4.0档沿用当前3.9分"),
        ("模拟真实体验分", "=4.6*35%+B15*20%+5*45%", "宝贝质量4.60、服务保障5.00及其他权重不变"),
    ]
    for r, values in enumerate(rows, 3):
        for c, value in enumerate(values, 1):
            ws.cell(r, c, value)
    for r in range(4, 17):
        ws.cell(r, 1).font = font(True)
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=GREY)
        ws.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 3).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 27
    for r in [7, 9, 10, 12, 13]:
        ws.cell(r, 2).number_format = "0.000"
    for r in [15, 16]:
        ws.cell(r, 2).number_format = "0.00"
    ws["B4"].number_format = "@"
    base_mean = df["物流到货时长/小时"].mean()
    excluded_mean = excluded["物流到货时长/小时"].mean()
    remaining = n - len(excluded)
    new_mean = (df["物流到货时长/小时"].sum() - excluded["物流到货时长/小时"].sum()) / remaining
    direction = "快于" if excluded_mean < base_mean else "慢于"
    tier = "4.0档或更高" if new_mean <= 58.66 else "未达4.0档"
    ws.merge_cells("A18:F18")
    ws["A18"] = f"结论：剔除{len(excluded)}笔，样本平均{direction}全店；剩余{remaining}笔平均{new_mean:.3f}小时，{tier}。详见上方公式。"
    ws["A18"].font = font(True, RED, 11)
    ws["A18"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[18].height = 38

    ws.cell(20, 1, f"被剔除的{len(excluded)}笔订单")
    ws.cell(20, 1).font = font(True, RED, 13)
    headers = ["订单编号", "商品ID", "物流线路", "快递公司", "到货时长/小时", "剔除原因"]
    for c, value in enumerate(headers, 1):
        ws.cell(21, c, value)
    for r, (_, row) in enumerate(excluded.iterrows(), 22):
        values = [str(row["订单编号"]), row["商品ID"], row["物流线路"], row["快递公司"], float(row["物流到货时长/小时"]), f"{city}收货＋指定商品"]
        for c, value in enumerate(values, 1):
            ws.cell(r, c, value)
        ws.cell(r, 1).number_format = "@"
        ws.cell(r, 2).number_format = "@"
        ws.cell(r, 5).number_format = "0.00"
        ws.row_dimensions[r].height = 24
    for col, width in enumerate([25, 20, 26, 16, 20, 28], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A21:F{21 + len(excluded)}"
    return ws


def add_exclusion_pickup_sheet(wb):
    """Recalculate every evening-pickup scenario after the temporary exclusion."""
    name = "临时剔除揽收测算"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = "999999"
    ws.merge_cells("A1:J1")
    ws["A1"] = "临时情景：剔除指定订单后的晚间揽收复算"
    ws["A1"].font = font(True, RED, 18)
    ws["A1"].fill = PatternFill("solid", fgColor=LIGHT)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    notes = [
        f"剔除口径：{os.environ['TMALL_LOGISTICS_EXCLUDE_CITY']}收货 × 商品 {os.environ['TMALL_LOGISTICS_EXCLUDE_PRODUCT_ID']}；仅作临时假设，不代表平台允许剔除。",
        "方案假设：沿用当前班次与候选新增班次；发货时间作为可交接时间，后续干线与签收按相同小时数提前。",
        "分母、受益订单、总节省和档位均在剔除后重新计算；实际效果仍需用真实揽收、发运和签收时间验证。",
    ]
    for row, text in enumerate(notes, 2):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        ws.cell(row, 1, text)
        ws.cell(row, 1).font = font(color="666666")
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 28

    scenario_value = os.environ.get(
        "TMALL_LOGISTICS_EXTRA_PICKUPS",
        "17:30,18:00,18:30,19:00,19:30,20:00,20:30,21:00,21:30,22:00,22:30,23:00",
    )
    times = scenario_value.split(",")
    if "18:00" not in times:
        raise ValueError("候选新增班次必须包含18:00")
    headers = [
        "新增揽收时间", "受益订单数", "总节省/小时", "剩余订单平均缩短/小时",
        "模拟平均到货/小时", "达到档位", "模拟物流速度分", "模拟真实体验分",
        "较18点少节省/小时", "执行观察",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(5, col, value)
    n = wb["订单明细"].max_row - 1
    best_row = 6 + times.index("18:00")
    arrival_fallback = float(os.environ.get("TMALL_LOGISTICS_ARRIVAL_SCORE", "3.9"))
    pickup_score = float(os.environ.get("TMALL_LOGISTICS_PICKUP_SCORE", "4.9"))
    for index, time_value in enumerate(times):
        row = 6 + index
        detail_col = get_column_letter(15 + index)
        source = f"'晚间揽收订单明细'!${detail_col}$2:${detail_col}${n + 1}"
        keep = (
            f"((订单明细!$I$2:$I${n + 1}<>'临时剔除测算'!$B$5)"
            f"+(订单明细!$R$2:$R${n + 1}<>'临时剔除测算'!$B$4)>0)"
        )
        ws.cell(row, 1, time_value)
        ws.cell(row, 2, f"=SUMPRODUCT(({source}>0)*{keep})")
        ws.cell(row, 3, f"=SUMPRODUCT({source}*{keep})")
        ws.cell(row, 4, f"=C{row}/'临时剔除测算'!$B$11")
        ws.cell(row, 5, f"='临时剔除测算'!$B$12-D{row}")
        arrival_score = (
            f"IF(E{row}<=43.99,5,IF(E{row}<=47.36,4.8,IF(E{row}<=49.91,4.6,"
            f"IF(E{row}<=56,4.2,IF(E{row}<=58.66,4,{arrival_fallback})))))"
        )
        ws.cell(row, 6, f'=IF(E{row}<=43.99,"5.0档",IF(E{row}<=47.36,"4.8档",IF(E{row}<=49.91,"4.6档",IF(E{row}<=56,"4.2档",IF(E{row}<=58.66,"4.0档","未达4.0档")))))')
        ws.cell(row, 7, f"={pickup_score}*60%+{arrival_score}*40%")
        ws.cell(row, 8, f"='临时剔除测算'!$B$16+(G{row}-'临时剔除测算'!$B$15)*20%")
        ws.cell(row, 9, f"=$C${best_row}-C{row}")
        ws.cell(row, 10, f'=IF(A{row}="18:00","原推荐时点复算",IF(B{row}<$B${best_row},"覆盖较少","覆盖相同，节省不增"))')
        for col in (3, 4, 5, 9):
            ws.cell(row, col).number_format = "0.000"
        for col in (7, 8):
            ws.cell(row, col).number_format = "0.00"
    last_row = 5 + len(times)
    style_header(ws, 5, 1, 10)
    style_data(ws, 6, last_row, 1, 10, numeric_cols=tuple(range(1, 10)), wrap_cols=(10,), row_height=28)
    for col in range(1, 11):
        ws.cell(best_row, col).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(best_row, col).font = font(True, RED)
    for col, width in enumerate([18, 16, 18, 24, 22, 18, 20, 20, 22, 28], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:J{last_row}"
    ws.merge_cells(start_row=last_row + 2, start_column=1, end_row=last_row + 2, end_column=10)
    ws.cell(last_row + 2, 1, "解读：比较‘模拟平均到货’与‘达到档位’，不要只看受益订单数；同步提前不代表实际提分保证。")
    ws.cell(last_row + 2, 1).font = font(True, RED)
    ws.cell(last_row + 2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[last_row + 2].height = 32
    return ws


def style_header(ws, row, start=1, end=None):
    end = end or ws.max_column
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = font(True, WHITE, 10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = Border(bottom=Side(style="medium", color=RED))
    ws.row_dimensions[row].height = 28


def style_data(ws, min_row, max_row, min_col, max_col, numeric_cols=(), wrap_cols=(), row_height=24, zebra=True):
    for row in range(min_row, max_row + 1):
        ws.row_dimensions[row].height = row_height
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row, col)
            cell.font = font()
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(
                horizontal="center" if col in numeric_cols else "left",
                vertical="center",
                wrap_text=col in wrap_cols,
            )
            if zebra and row % 2 == 0 and cell.fill.fill_type is None:
                cell.fill = PatternFill("solid", fgColor="FAFAFA")


def polish_workbook(wb):
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "landscape"
        ws.sheet_view.zoomScale = 90
        ws.sheet_properties.tabColor = BLUE

    summary = wb["分析摘要"]
    summary.sheet_properties.tabColor = RED
    summary.freeze_panes = "A3"
    for col, width in {"A": 50, "B": 22, "C": 58, "D": 18, "E": 14, "F": 14}.items():
        summary.column_dimensions[col].width = width
    for r in [3, 4, 5, 6, 7, 8, 9, 10, 11, 12]:
        summary.cell(r, 1).font = font(True)
        summary.cell(r, 2).font = font()
        summary.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")
        summary.row_dimensions[r].height = 25
    style_header(summary, 24, 1, 4)
    style_data(summary, 25, 30, 1, 4, numeric_cols=(2, 3), row_height=28)
    style_header(summary, 33, 1, 3)
    style_data(summary, 34, 40, 1, 3, numeric_cols=(2,), wrap_cols=(3,), row_height=32)
    for r in range(34, 41):
        summary.cell(r, 1).font = font(True)
    for r in [16, 17, 18, 19, 20, 42, 43]:
        summary.row_dimensions[r].height = 34
        summary.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")

    evening = wb["晚间揽收分析"]
    evening.sheet_properties.tabColor = RED
    extra_times = os.environ.get("TMALL_LOGISTICS_EXTRA_PICKUPS", "17:30,18:00,18:30,19:00,19:30,20:00,20:30,21:00,21:30,22:00,22:30,23:00").split(",")
    best_row = 7 + extra_times.index("18:00")
    last_scenario_row = 6 + len(extra_times)
    for col, width in enumerate([18, 14, 16, 22, 16, 20, 20, 16, 22, 26], 1):
        evening.column_dimensions[get_column_letter(col)].width = width
    style_header(evening, 6, 1, 10)
    style_data(evening, 7, last_scenario_row, 1, 10, numeric_cols=(2, 3, 4, 5, 6, 7, 8, 9), row_height=26)
    for col in range(1, 11):
        evening.cell(best_row, col).fill = PatternFill("solid", fgColor=LIGHT)
        evening.cell(best_row, col).font = font(True, RED)
    for cell in evening[22]:
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()
    for col, value in enumerate(["实际兑现比例", "全店缩短/小时", "模拟到货/小时", "达到档位"], 1):
        evening.cell(22, col, value)
    style_header(evening, 22, 1, 4)
    style_data(evening, 23, 27, 1, 4, numeric_cols=(1, 2, 3, 4), row_height=27, zebra=False)
    for r in range(23, 28):
        evening.cell(r, 5).value = None
    evening.merge_cells("A29:J29")
    evening["A29"] = "说明：兑现比例用于检验实际效果低于‘同步提前’情景时的档位变化；0%表示仅扫描提前、后续完全未提前。"
    evening["A29"].alignment = Alignment(wrap_text=True, vertical="center")
    evening["A29"].font = font(color="666666")
    evening.row_dimensions[29].height = 32
    for r in [2, 3, 4]:
        evening.row_dimensions[r].height = 28
        evening.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")

    scenarios = wb["情景测算"]
    for col, width in enumerate([48, 24, 20, 18, 58], 1):
        scenarios.column_dimensions[get_column_letter(col)].width = width
    style_header(scenarios, 1, 1, 5)
    style_data(scenarios, 2, 8, 1, 5, numeric_cols=(2, 3, 4), wrap_cols=(5,), row_height=32)

    cities = wb["城市分析"]
    style_header(cities, 1, 1, 13)
    style_data(cities, 2, cities.max_row, 1, 13, numeric_cols=tuple(range(2, 14)), row_height=25)
    cities.freeze_panes = "B2"

    routes = wb["线路分析"]
    style_header(routes, 1, 1, 10)
    style_data(routes, 2, routes.max_row, 1, 10, numeric_cols=tuple(range(3, 11)), row_height=25)
    routes.freeze_panes = "C2"

    orders = wb["订单明细"]
    style_header(orders, 1, 1, 18)
    style_data(orders, 2, orders.max_row, 1, 18, numeric_cols=tuple(range(6, 19)), wrap_cols=(2, 3), row_height=32, zebra=False)
    orders.freeze_panes = "D2"
    orders.sheet_view.zoomScale = 75

    pickup = wb["晚间揽收订单明细"]
    style_header(pickup, 1, 1, pickup.max_column)
    style_data(pickup, 2, pickup.max_row, 1, pickup.max_column, numeric_cols=tuple(range(7, pickup.max_column + 1)), wrap_cols=(4,), row_height=30, zebra=False)
    pickup.freeze_panes = "D2"
    pickup.sheet_view.zoomScale = 75

    sources = wb["计分规则与来源"]
    for col, width in enumerate([28, 72, 38, 58], 1):
        sources.column_dimensions[get_column_letter(col)].width = width
    for row in [2, 13, 23]:
        style_header(sources, row, 1, 4 if row != 13 else 2)
    style_data(sources, 3, 11, 1, 4, wrap_cols=(2, 3, 4), row_height=38, zebra=False)
    style_data(sources, 14, 19, 1, 2, numeric_cols=(2,), row_height=25, zebra=False)
    style_data(sources, 24, sources.max_row, 1, 4, wrap_cols=(2, 3, 4), row_height=40, zebra=False)
    sources.freeze_panes = "A3"

    if "临时剔除测算" in wb.sheetnames:
        temporary = wb["临时剔除测算"]
        temporary.sheet_properties.tabColor = "999999"
        style_header(temporary, 3, 1, 3)
        style_header(temporary, 21, 1, 6)
        style_data(temporary, 22, temporary.max_row, 1, 6, numeric_cols=(2, 5), row_height=26)
        for r in range(22, temporary.max_row + 1):
            temporary.cell(r, 1).number_format = "@"
            temporary.cell(r, 2).number_format = "@"
        for r in range(4, 17):
            temporary.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")

    desired_order = ["分析摘要", "晚间揽收分析", "情景测算", "临时剔除测算", "临时剔除揽收测算", "城市分析", "线路分析", "订单明细", "晚间揽收订单明细", "计分规则与来源"]
    wb._sheets = [wb[name] for name in desired_order if name in wb.sheetnames]
    wb.active = 0


def main():
    wb = load_workbook(BOOK)
    if os.environ.get("TMALL_LOGISTICS_EXCLUDE_PRODUCT_ID") and os.environ.get("TMALL_LOGISTICS_EXCLUDE_CITY"):
        add_exclusion_sheet(wb)
        add_exclusion_pickup_sheet(wb)
    polish_workbook(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(BOOK)
    print(BOOK)


if __name__ == "__main__":
    main()
